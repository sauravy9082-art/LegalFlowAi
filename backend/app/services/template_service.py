import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from typing import Any, Optional
from app.models.schemas import TemplateGenerateRequest
from app.core.config import settings

DISCLAIMER_TEXT = "⚠️ AI-GENERATED DRAFT — REVIEW BY A QUALIFIED LEGAL PROFESSIONAL BEFORE SIGNING."

def format_governing_law(law_input: Optional[str]) -> str:
    """Format governing law string cleanly without duplicate 'laws of' prefixes."""
    if not law_input or not law_input.strip():
        return "the applicable laws of the jurisdiction"
    clean = law_input.strip()
    clean_lower = clean.lower()
    if clean_lower.startswith("laws of ") or clean_lower.startswith("the laws of "):
        return clean
    elif clean_lower.startswith("state of ") or clean_lower.startswith("the state of "):
        return f"the laws of the {clean}"
    else:
        return f"the laws of {clean}"

class TemplateService:
    """
    Form-Driven Contract Generator Engine (.DOCX) (Rule 26).
    Generates template-specific customized agreements with mandatory legal disclaimers.
    """

    async def generate_contract(self, req: TemplateGenerateRequest) -> str:
        os.makedirs(os.path.join(settings.STORAGE_PATH, "templates"), exist_ok=True)

        # Extract values from form_data if provided
        f_data = req.form_data or {}
        
        party_a = (
            f_data.get("party_a_name") or
            f_data.get("disclosing_party") or
            f_data.get("landlord_name") or
            f_data.get("client_name") or
            req.party_a_name or
            "Party A"
        )
        party_b = (
            f_data.get("party_b_name") or
            f_data.get("receiving_party") or
            f_data.get("tenant_name") or
            f_data.get("freelancer_name") or
            req.party_b_name or
            "Party B"
        )
        effective_date = (
            f_data.get("effective_date") or
            f_data.get("lease_start") or
            f_data.get("agreement_date") or
            req.effective_date or
            "2026-08-29"
        )
        duration_months = (
            f_data.get("duration_months") or
            f_data.get("term_months") or
            req.duration_months or
            24
        )
        payment_terms = (
            f_data.get("payment_terms") or
            f_data.get("payment_amount") or
            f_data.get("monthly_rent") or
            req.payment_terms or
            "Net 30 days upon monthly invoice receipt"
        )
        governing_law_raw = (
            f_data.get("governing_law") or
            f_data.get("jurisdiction") or
            req.governing_law or
            "Laws of India"
        )
        governing_law_str = format_governing_law(governing_law_raw)

        # Determine exact template key
        t_id = (req.template_id or "").lower()
        t_cat = (req.template_type or "").upper()
        
        filename_safe_a = party_a.replace(" ", "_").replace("/", "_")
        filename = f"Generated_{req.template_type}_{filename_safe_a}.docx"
        file_path = os.path.join(settings.STORAGE_PATH, "templates", filename)

        doc = docx.Document()

        # Page margins
        for section in doc.sections:
            section.top_margin = Inches(1)
            section.bottom_margin = Inches(1)
            section.left_margin = Inches(1)
            section.right_margin = Inches(1)

        # Header Disclaimer Banner
        header_p = doc.add_paragraph()
        header_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        h_run = header_p.add_run(DISCLAIMER_TEXT)
        h_run.bold = True
        h_run.font.size = Pt(9.5)
        h_run.font.color.rgb = RGBColor(180, 0, 0)

        # Build Template-Specific Content
        if "nda" in t_id or "confidentiality" in t_cat or t_cat == "NDA":
            self._build_nda_document(doc, req, party_a, party_b, effective_date, duration_months, governing_law_str)
        elif "lease" in t_id or "real_estate" in t_cat or "property" in t_id:
            self._build_lease_document(doc, req, party_a, party_b, effective_date, payment_terms, governing_law_str)
        elif "service" in t_id or "msa" in t_id or "services" in t_cat or t_cat == "MSA":
            self._build_msa_document(doc, req, party_a, party_b, effective_date, payment_terms, duration_months, governing_law_str)
        else:
            self._build_general_agreement_document(doc, req, party_a, party_b, effective_date, payment_terms, duration_months, governing_law_str)

        # Signatures Section
        doc.add_heading("Signatures & Execution", level=2)
        table = doc.add_table(rows=3, cols=2)
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = f"For Disclosing / First Party ({party_a}):"
        hdr_cells[1].text = f"For Receiving / Second Party ({party_b}):"

        r1 = table.rows[1].cells
        r1[0].text = "\nAuthorized Signature: ____________________\nDate: _______________"
        r1[1].text = "\nAuthorized Signature: ____________________\nDate: _______________"

        # Footer Disclaimer
        footer_p = doc.add_paragraph()
        footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        f_run = footer_p.add_run(f"\n\n{DISCLAIMER_TEXT}")
        f_run.italic = True
        f_run.font.size = Pt(8.5)
        f_run.font.color.rgb = RGBColor(120, 120, 120)

        doc.save(file_path)
        return file_path

    def _build_nda_document(self, doc: docx.Document, req: TemplateGenerateRequest, party_a: str, party_b: str, effective_date: str, duration_months: Any, governing_law_str: str):
        """Build Mutual Non-Disclosure Agreement (NDA)"""
        title_p = doc.add_paragraph()
        title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        t_run = title_p.add_run("\nMUTUAL NON-DISCLOSURE AGREEMENT (NDA)\n")
        t_run.bold = True
        t_run.font.size = Pt(16)

        doc.add_paragraph(
            f"This Mutual Non-Disclosure Agreement (\"Agreement\") is entered into and made effective as of "
            f"{effective_date}, by and between {party_a} (\"Disclosing Party / Party A\") and {party_b} (\"Receiving Party / Party B\")."
        )

        doc.add_heading("1. Purpose", level=2)
        doc.add_paragraph(
            "The parties wish to explore potential commercial discussions and technical evaluation (the \"Purpose\"). "
            "In connection with the Purpose, each party may disclose to the other party proprietary, technical, financial, and trade secret information."
        )

        doc.add_heading("2. Definition of Confidential Information", level=2)
        doc.add_paragraph(
            "\"Confidential Information\" shall include all non-public information disclosed by one party to the other, "
            "whether orally, visually, or in writing, including business strategies, customer data, software code, financial data, and technical specifications."
        )

        doc.add_heading("3. Obligations & Non-Disclosure", level=2)
        doc.add_paragraph(
            "Each party agrees to maintain all Confidential Information in strict confidence and shall not disclose such information "
            "to any third party without prior written authorization from the disclosing party. Confidential Information shall be used solely for the Purpose."
        )

        doc.add_heading("4. Term & Duration", level=2)
        doc.add_paragraph(
            f"This Agreement and all confidentiality obligations hereunder shall remain in effect for a term of {duration_months} month(s) "
            f"from the Effective Date of {effective_date}."
        )

        doc.add_heading("5. Governing Law & Jurisdiction", level=2)
        doc.add_paragraph(
            f"This Agreement shall be governed by, construed, and enforced in accordance with {governing_law_str}."
        )

    def _build_msa_document(self, doc: docx.Document, req: TemplateGenerateRequest, party_a: str, party_b: str, effective_date: str, payment_terms: str, duration_months: Any, governing_law_str: str):
        """Build Master Services Agreement (MSA)"""
        title_p = doc.add_paragraph()
        title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        t_run = title_p.add_run("\nMASTER SERVICES AGREEMENT (MSA)\n")
        t_run.bold = True
        t_run.font.size = Pt(16)

        doc.add_paragraph(
            f"This Master Services Agreement (\"Agreement\") is entered into and made effective as of "
            f"{effective_date}, by and between {party_a} (\"Client\") and {party_b} (\"Service Provider\")."
        )

        doc.add_heading("1. Services & Statements of Work", level=2)
        doc.add_paragraph(
            f"Service Provider agrees to perform technical services, software development, or consulting deliverables "
            f"as specified in Statements of Work (SOW) executed pursuant to this Agreement for an initial term of {duration_months} month(s)."
        )

        doc.add_heading("2. Compensation & Payment Structure", level=2)
        doc.add_paragraph(
            f"Client agrees to pay Service Provider in accordance with the following terms: {payment_terms}. "
            f"Invoices shall be rendered monthly and paid according to the agreed schedule."
        )

        doc.add_heading("3. Intellectual Property Assignment", level=2)
        doc.add_paragraph(
            "Upon receipt of full payment by Client, Service Provider hereby assigns to Client all right, title, and interest "
            "in and to all custom deliverables created specifically for Client under this Agreement."
        )

        doc.add_heading("4. Confidentiality & Non-Disclosure", level=2)
        doc.add_paragraph(
            "Both parties agree to hold all technical data, business systems, and customer records in strict confidence during and after the term."
        )

        doc.add_heading("5. Termination & Notice", level=2)
        doc.add_paragraph(
            "Either party may terminate this Agreement for convenience upon thirty (30) days' prior written notice. "
            "In the event of a material breach, the non-breaching party may terminate immediately after a 15-day opportunity to cure."
        )

        doc.add_heading("6. Governing Law & Jurisdiction", level=2)
        doc.add_paragraph(
            f"This Agreement shall be governed by, construed, and enforced in accordance with {governing_law_str}."
        )

    def _build_lease_document(self, doc: docx.Document, req: TemplateGenerateRequest, party_a: str, party_b: str, effective_date: str, payment_terms: str, governing_law_str: str):
        """Build Commercial Lease & Property Agreement"""
        title_p = doc.add_paragraph()
        title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        t_run = title_p.add_run("\nCOMMERCIAL LEASE & PROPERTY AGREEMENT\n")
        t_run.bold = True
        t_run.font.size = Pt(16)

        doc.add_paragraph(
            f"This Commercial Lease Agreement (\"Lease\") is entered into and made effective as of "
            f"{effective_date}, by and between {party_a} (\"Lessor / Landlord\") and {party_b} (\"Lessee / Tenant\")."
        )

        doc.add_heading("1. Demised Premises & Lease Term", level=2)
        doc.add_paragraph(
            f"Lessor hereby leases to Lessee commercial office and premises starting on the Lease Start Date of {effective_date}. "
            f"Lessee agrees to occupy the premises in compliance with all zoning and municipal laws."
        )

        doc.add_heading("2. Rent & Payment Terms", level=2)
        doc.add_paragraph(
            f"Lessee agrees to pay Lessor monthly rent in accordance with the agreed schedule: {payment_terms}. "
            f"Rent is payable on or before the due date specified in each billing period."
        )

        doc.add_heading("3. Security Deposit & Maintenance Covenants", level=2)
        doc.add_paragraph(
            "Lessee agrees to maintain the premises in good working order and repair. "
            "Any security deposit paid by Lessee shall be refunded upon lease expiration, less documented repairs for tenant damages."
        )

        doc.add_heading("4. Default & Eviction Remedy", level=2)
        doc.add_paragraph(
            "In the event of non-payment or breach, Lessor shall provide written notice and a 14-day cure period prior to exercising legal default remedies."
        )

        doc.add_heading("5. Governing Law & Jurisdiction", level=2)
        doc.add_paragraph(
            f"This Lease Agreement shall be governed by, construed, and enforced in accordance with {governing_law_str}."
        )

    def _build_general_agreement_document(self, doc: docx.Document, req: TemplateGenerateRequest, party_a: str, party_b: str, effective_date: str, payment_terms: str, duration_months: Any, governing_law_str: str):
        """Build General Commercial Agreement"""
        title_p = doc.add_paragraph()
        title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        t_run = title_p.add_run(f"\n{req.template_type.upper()} AGREEMENT\n")
        t_run.bold = True
        t_run.font.size = Pt(16)

        doc.add_paragraph(
            f"This {req.template_type} Agreement (\"Agreement\") is entered into and made effective as of "
            f"{effective_date}, by and between {party_a} (\"Party A\") and {party_b} (\"Party B\")."
        )

        doc.add_heading("1. Purpose & Scope", level=2)
        doc.add_paragraph(
            f"The parties agree to engage under the terms set forth herein for a duration of {duration_months} month(s)."
        )

        doc.add_heading("2. Payment & Terms", level=2)
        doc.add_paragraph(f"Payment Schedule & Structure: {payment_terms}")

        doc.add_heading("3. Confidentiality", level=2)
        doc.add_paragraph("Each party agrees to maintain in strict confidence all proprietary technical, financial, and business information.")

        doc.add_heading("4. Governing Law & Jurisdiction", level=2)
        doc.add_paragraph(f"This Agreement shall be governed by, construed, and enforced in accordance with {governing_law_str}.")
